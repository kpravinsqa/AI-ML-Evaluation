"""
Model evaluation workbook for model-risk / release sign-off.

    python -m reports_tool.model_eval_report    -> reports/model_evaluation_report.xlsx
"""
from __future__ import annotations

import json
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sklearn.metrics import brier_score_loss, f1_score, precision_score, recall_score, roc_auc_score

from data.generate_data import TARGET, production_batches, splits
from src.features import MODEL_INPUTS
from src.metrics import disparate_impact_ratio, equal_opportunity_difference, expected_calibration_error, feature_drift_report
from src.train import MODEL_VERSION, build

ROOT = Path(__file__).resolve().parent.parent
HEAD, OK, BAD = PatternFill("solid", fgColor="1F4E79"), PatternFill("solid", fgColor="C6EFCE"), PatternFill("solid", fgColor="FFC7CE")


def add_sheet(wb, title, header, rows, status_col=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
    for r in rows:
        ws.append(list(r))
        if status_col:
            ws.cell(ws.max_row, status_col).fill = OK if r[status_col - 1] in ("PASS", "STABLE") else BAD
    for col in ws.columns:
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(len(str(c.value or "")) for c in col) + 2, 50)


def main(out: Path = ROOT / "reports" / "model_evaluation_report.xlsx") -> Path:
    gates = json.loads((ROOT / "config" / "quality_gates.json").read_text())
    train, test = splits()
    model = build("baseline").fit(train[MODEL_INPUTS], train[TARGET])
    challenger = build("challenger").fit(train[MODEL_INPUTS], train[TARGET])
    p = model.predict_proba(test[MODEL_INPUTS])[:, 1]
    y = test[TARGET].values
    approved = (p < gates["inference"]["decision_threshold_approve"]).astype(int)
    yhat = (p >= 0.3).astype(int)

    def gate(name, value, op, limit):
        ok = value >= limit if op == ">=" else value <= limit
        return (name, round(float(value), 4), f"{op} {limit}", "PASS" if ok else "FAIL")

    rows = [
        gate("ROC-AUC", roc_auc_score(y, p), ">=", gates["performance"]["roc_auc_min"]),
        gate("Recall (defaulters @0.30)", recall_score(y, yhat), ">=", gates["performance"]["recall_default_min"]),
        gate("Precision (defaulters @0.30)", precision_score(y, yhat), ">=", gates["performance"]["precision_default_min"]),
        gate("F1 (@0.30)", f1_score(y, yhat), ">=", 0.5),
        gate("Expected calibration error", expected_calibration_error(y, p), "<=", gates["calibration"]["ece_max"]),
        gate("Brier score", brier_score_loss(y, p), "<=", gates["calibration"]["brier_max"]),
        gate("Disparate impact (gender)", disparate_impact_ratio(approved, test["gender"].values), ">=", gates["fairness"]["disparate_impact_min"]),
        gate("Disparate impact (age band)", disparate_impact_ratio(approved, test["age_band"].values), ">=", gates["fairness"]["disparate_impact_min"]),
        gate("Equal opportunity diff (gender)", equal_opportunity_difference(y, approved, test["gender"].values), "<=", gates["fairness"]["equal_opportunity_diff_max"]),
    ]
    wb = Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "Summary", ["Item", "Value"], [
        ("Model", MODEL_VERSION), ("Champion", "Logistic regression"), ("Challenger", "Gradient boosting"),
        ("Champion AUC", round(roc_auc_score(y, p), 4)),
        ("Challenger AUC", round(roc_auc_score(y, challenger.predict_proba(test[MODEL_INPUTS])[:, 1]), 4)),
        ("Hold-out size", len(test)), ("Observed default rate", round(float(y.mean()), 4)),
        ("Gates passed", f"{sum(r[3] == 'PASS' for r in rows)}/{len(rows)}")])
    add_sheet(wb, "Quality gates", ["Metric", "Value", "Gate", "Status"], rows, status_col=4)
    seg = []
    for attr in ("gender", "age_band", "home_ownership"):
        for v in sorted(test[attr].unique()):
            m = (test[attr] == v).values
            seg.append((attr, v, int(m.sum()), round(float(y[m].mean()), 3), round(float(p[m].mean()), 3),
                        round(roc_auc_score(y[m], p[m]), 3), round(float(approved[m].mean()), 3)))
    add_sheet(wb, "Segments", ["Attribute", "Group", "N", "Observed DR", "Mean PD", "AUC", "Approval rate"], seg)
    stable, drifted = production_batches()
    feats = ["debt_to_income", "credit_utilisation", "missed_payments_12m", "annual_income", "loan_amount"]
    d = []
    for name, batch in (("stable", stable), ("downturn", drifted)):
        for r in feature_drift_report(train, batch, feats).to_dict("records"):
            d.append((name, r["feature"], r["psi"], round(r["ks_pvalue"], 6), r["status"]))
    add_sheet(wb, "Drift", ["Batch", "Feature", "PSI", "KS p-value", "Status"], d, status_col=5)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


if __name__ == "__main__":
    print("Report written to", main())
